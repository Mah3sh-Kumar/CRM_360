import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from datetime import date

from app.main import app
from app.core.database import Base, get_db
from app.models.company import Company
from app.models.employee import Employee
from app.models.lead import Lead
from app.models.customer import Customer
from app.core.security import get_password_hash

SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"
engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def override_get_db():
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()


client = TestClient(app)

@pytest.fixture(scope="module", autouse=True)
def setup_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    
    # Create companies
    c1 = Company(id=1, company_code="C1", name="Company 1", email="c1@test.com", status="Active")
    db.add(c1)
    db.commit()
    
    # Create employees
    d = date(2023, 1, 1)
    u_sa = Employee(id=1, employee_code="SA", name="SA", email="sa@test.com", mobile="1", hashed_password=get_password_hash("test"), role="Super Admin", company_id=None, designation="SA", date_of_joining=d, status="Active")
    db.add(u_sa)
    db.commit()
    
    # E1
    u_e1 = Employee(id=2, employee_code="E1", name="E1", email="e1@test.com", mobile="2", hashed_password=get_password_hash("test"), role="Employee", company_id=1, designation="EMP", date_of_joining=d, status="Active")
    db.add(u_e1)
    db.commit()

    # Create dummy data
    lead1 = Lead(lead_name="L1", contact_person="C1", mobile_number="1", lead_source="Web", status="New", company_id=1, assigned_employee_id=1)
    db.add(lead1)
    db.commit()
    db.close()

@pytest.fixture
def sa_token():
    response = client.post(
        "/api/auth/login",
        json={"email": "sa@test.com", "password": "test"}
    )
    return response.json()["access_token"]

@pytest.fixture
def sa_headers(sa_token):
    return {"Authorization": f"Bearer {sa_token}"}

@pytest.fixture
def e1_token():
    response = client.post(
        "/api/auth/login",
        json={"email": "e1@test.com", "password": "test"}
    )
    return response.json()["access_token"]

@pytest.fixture
def e1_headers(e1_token):
    return {"Authorization": f"Bearer {e1_token}"}

# List of all newly added endpoints to test authorization and basic data structure
ENDPOINTS = [
    "/api/reports/leads/source-analysis",
    "/api/reports/leads/conversion",
    "/api/reports/leads/lost-analysis",
    "/api/reports/leads/employee-performance",
    "/api/reports/customers/ownership",
    "/api/reports/customers/status",
    "/api/reports/customers/industry-analysis",
    "/api/reports/opportunities/revenue-forecast",
    "/api/reports/opportunities/won",
    "/api/reports/opportunities/lost",
    "/api/reports/opportunities/aging",
    "/api/reports/activities/summary",
    "/api/reports/activities/employee",
    "/api/reports/activities/customer-interaction",
    "/api/reports/activities/pending",
    "/api/reports/tasks/summary",
    "/api/reports/tasks/pending",
    "/api/reports/tasks/overdue",
    "/api/reports/tasks/completion",
    "/api/reports/meetings/summary",
    "/api/reports/meetings/customer",
    "/api/reports/meetings/employee",
    "/api/reports/meetings/outcome",
    "/api/reports/visits/daily",
    "/api/reports/visits/customer",
    "/api/reports/visits/employee",
    "/api/reports/visits/outcome",
    "/api/reports/attendance/daily",
    "/api/reports/attendance/working-hours",
    "/api/reports/attendance/monthly",
    "/api/reports/geo/tracking",
]

@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_endpoint_authorization_missing(endpoint):
    response = client.get(endpoint)
    assert response.status_code == 401

@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_endpoint_sa_success(endpoint, sa_headers):
    response = client.get(f"{endpoint}?format=json", headers=sa_headers)
    assert response.status_code == 200
    data = response.json()
    assert "data" in data

@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_endpoint_employee_success(endpoint, e1_headers):
    # Tests that an employee can fetch it and it respects isolation (doesn't throw 500)
    response = client.get(f"{endpoint}?format=json", headers=e1_headers)
    assert response.status_code == 200
    data = response.json()
    assert "data" in data

# Specific data correctness tests

def test_lead_conversion_calculation(sa_headers):
    response = client.get("/api/reports/leads/conversion?format=json", headers=sa_headers)
    assert response.status_code == 200
    data = response.json()
    # verify conversion calculation
    for row in data["data"]:
        assert "conversion_rate_percent" in row
        total = row["total_leads"]
        won = row["won_leads"]
        expected = round((won / total * 100), 2) if total > 0 else 0
        assert row["conversion_rate_percent"] == expected

def test_task_completion_calculation(sa_headers):
    response = client.get("/api/reports/tasks/completion?format=json", headers=sa_headers)
    assert response.status_code == 200
    data = response.json()
    for row in data["data"]:
        assert "completion_rate" in row
        total = row["total_tasks"]
        comp = row["completed_tasks"]
        expected = round((comp / total * 100), 2) if total > 0 else 0
        assert row["completion_rate"] == expected

def test_revenue_forecast(sa_headers):
    response = client.get("/api/reports/opportunities/revenue-forecast?format=json", headers=sa_headers)
    assert response.status_code == 200




def test_csv_export_parsing(sa_headers):
    response = client.get("/api/reports/leads/summary?format=csv", headers=sa_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    
    import csv, io
    reader = csv.DictReader(io.StringIO(response.text))
    rows = list(reader)
    assert len(reader.fieldnames) > 0
    if len(rows) > 0:
        assert "Status" in rows[0]

def test_xlsx_export_parsing(sa_headers):
    response = client.get("/api/reports/leads/summary?format=xlsx", headers=sa_headers)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    
    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.max_row >= 1
    assert ws.cell(row=1, column=1).value == "Status"

def test_pdf_export_parsing(sa_headers):
    response = client.get("/api/reports/leads/summary?format=pdf", headers=sa_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    
    from pypdf import PdfReader
    import io
    reader = PdfReader(io.BytesIO(response.content))
    assert len(reader.pages) > 0
    text = reader.pages[0].extract_text()
    assert "Lead Summary" in text or "Status" in text

@pytest.fixture(autouse=True)
def apply_override():
    app.dependency_overrides[get_db] = override_get_db
